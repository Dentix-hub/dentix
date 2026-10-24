import io
import pytest
from decimal import Decimal
from datetime import datetime, timezone
from openpyxl import load_workbook
from httpx import AsyncClient, ASGITransport

from backend import models, schemas
from backend.main import app
from backend.auth import create_access_token


@pytest.mark.asyncio
async def test_excel_export_permission_denied_for_assistant(async_db_session):
    """Assistant without SYSTEM_CONFIG permission must receive 403."""
    tenant = models.Tenant(name="Clinic Test", subscription_status="active", is_active=True)
    async_db_session.add(tenant)
    await async_db_session.flush()

    assistant = models.User(
        username="assistant_user",
        email="assistant@example.com",
        role="assistant",
        tenant_id=tenant.id,
        hashed_password="hash",
    )
    async_db_session.add(assistant)
    await async_db_session.commit()

    token = create_access_token(data={"sub": assistant.username, "tenant_id": tenant.id, "role": assistant.role})

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        response = await ac.get(
            "/api/v1/settings/backup/export-xlsx",
            headers={"Authorization": f"Bearer {token}"},
        )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_excel_export_tenant_isolation_and_formula_sanitization(async_db_session):
    """
    Export must be tenant-isolated (Tenant A sees only its own rows).
    Formula injection in string cells must be sanitized with leading single quote.
    """
    now = datetime.now(timezone.utc).replace(tzinfo=None)

    # 1. Create Tenant A & Tenant B
    tenant_a = models.Tenant(name="Tenant Alpha", subscription_status="active", is_active=True)
    tenant_b = models.Tenant(name="Tenant Beta", subscription_status="active", is_active=True)
    async_db_session.add_all([tenant_a, tenant_b])
    await async_db_session.flush()

    admin_a = models.User(
        username="admin_a",
        email="admin_a@example.com",
        role="admin",
        tenant_id=tenant_a.id,
        hashed_password="hash",
    )
    admin_b = models.User(
        username="admin_b",
        email="admin_b@example.com",
        role="admin",
        tenant_id=tenant_b.id,
        hashed_password="hash",
    )
    async_db_session.add_all([admin_a, admin_b])
    await async_db_session.flush()

    # 2. Seed data for Tenant A with formula injection attack strings
    patient_a = models.Patient(
        tenant_id=tenant_a.id,
        name="=cmd|' /C calc'!A0",  # Formula injection payload
        age=30,
        phone="07701111111",
        medical_history="+Allergy to Penicillin",  # Starts with +
        notes="@SUM(1,2)",  # Starts with @
        created_at=now,
    )
    appt_a = models.Appointment(
        tenant_id=tenant_a.id,
        patient=patient_a,
        doctor_id=admin_a.id,
        date_time=now,
        status="Scheduled",
        duration_minutes=30,
        notes="-Negative note",  # Starts with -
    )
    treat_a = models.Treatment(
        tenant_id=tenant_a.id,
        patient=patient_a,
        doctor_id=admin_a.id,
        procedure="Composite Filling",
        cost=Decimal("120.00"),
        discount=Decimal("20.00"),
        date=now,
        status="Done",
        notes="\tTab prefixed note",  # Starts with \t
    )
    pay_a = models.Payment(
        tenant_id=tenant_a.id,
        patient=patient_a,
        doctor_id=admin_a.id,
        amount=Decimal("100.00"),
        date=now,
        notes="\rCR prefixed note",  # Starts with \r
    )
    proc_a = models.Procedure(
        tenant_id=tenant_a.id,
        name="=DANGEROUS_PROC()",
        price=Decimal("120.00"),
    )
    plan_a = models.ClinicalTreatmentPlan(
        tenant_id=tenant_a.id,
        patient=patient_a,
        created_by_user_id=admin_a.id,
        title="+Special Plan",
        status="active",
        created_at=now,
    )

    # 3. Seed data for Tenant B (must NOT appear in Tenant A's export)
    patient_b = models.Patient(
        tenant_id=tenant_b.id,
        name="Tenant B Secret Patient",
        age=45,
        phone="07709999999",
        medical_history="Secret",
        notes="Secret B notes",
        created_at=now,
    )
    proc_b = models.Procedure(
        tenant_id=tenant_b.id,
        name="Tenant B Secret Procedure",
        price=Decimal("500.00"),
    )

    async_db_session.add_all([
        patient_a, appt_a, treat_a, pay_a, proc_a, plan_a,
        patient_b, proc_b,
    ])
    await async_db_session.commit()

    # 4. Request export as Admin A
    token_a = create_access_token(data={"sub": admin_a.username, "tenant_id": tenant_a.id, "role": admin_a.role})
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        response = await ac.get(
            "/api/v1/settings/backup/export-xlsx",
            headers={"Authorization": f"Bearer {token_a}"},
        )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "dentix-backup-tenant_alpha" in response.headers.get("content-disposition", "")

    # 5. Load and inspect workbook
    wb = load_workbook(io.BytesIO(response.content))
    expected_sheets = [
        "Patients",
        "Appointments",
        "Treatments",
        "Payments",
        "Procedures",
        "Inventory Items",
        "Lab Orders",
        "Treatment Plans",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Sheet {s} missing from workbook"

    # Inspect Patients sheet
    ws_patients = wb["Patients"]
    patient_rows = list(ws_patients.iter_rows(values_only=True))
    assert len(patient_rows) == 2  # 1 header + 1 row for Tenant A
    # Verify Tenant B patient is NOT present
    for row in patient_rows:
        for cell in row:
            assert cell != "Tenant B Secret Patient", "Tenant B data leaked to Tenant A!"

    # Verify formula sanitization on patient name: starts with '
    patient_name_cell = patient_rows[1][1]
    assert patient_name_cell.startswith("'="), f"Expected formula sanitization prefix, got: {patient_name_cell}"

    # Verify medical history: starts with '+
    med_history_cell = patient_rows[1][6]
    assert med_history_cell.startswith("'+"), f"Expected formula sanitization prefix, got: {med_history_cell}"

    # Inspect Treatments sheet
    ws_treatments = wb["Treatments"]
    treatment_rows = list(ws_treatments.iter_rows(values_only=True))
    assert len(treatment_rows) == 2
    # Verify numeric money values
    cost_val = treatment_rows[1][7]
    assert isinstance(cost_val, (int, float))
    assert round(float(cost_val), 2) == 120.00
