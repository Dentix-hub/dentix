import io
import asyncio
import logging
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import joinedload
from sqlalchemy.ext.asyncio import AsyncSession

from backend import models

logger = logging.getLogger("dentix.excel_export")

ROW_LIMIT = 200_000
TRUNCATION_MARKER = ["[TRUNCATED: Sheet row limit of 200,000 reached]"]


def sanitize_cell_value(val: Any) -> Any:
    """
    Sanitize and format spreadsheet cell values.
    - None -> ""
    - Booleans -> bool
    - Decimal -> float quantized to 2 decimal places (HALF_UP)
    - int / float -> numeric value preserved
    - datetime / date -> YYYY-MM-DD HH:MM or YYYY-MM-DD
    - Strings -> preserved as string without prepending quote (formula protection via cell.data_type = 's')
    """
    if val is None:
        return ""
    if isinstance(val, bool):
        return val
    if isinstance(val, Decimal):
        return float(val.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")

    return str(val)


def _build_workbook_sync(sheets_data: dict[str, tuple[list[str], list[list[Any]]]]) -> io.BytesIO:
    """
    Synchronous openpyxl workbook generator executed in asyncio.to_thread.
    """
    wb = Workbook()
    first = True
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for sheet_name, (headers, rows) in sheets_data.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=sheet_name)

        # Header row
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.data_type = "s"

        # Data rows
        for row in rows:
            clean_row = [sanitize_cell_value(v) for v in row]
            ws.append(clean_row)
            for col_idx, cell_val in enumerate(clean_row, 1):
                if isinstance(cell_val, str):
                    ws.cell(row=ws.max_row, column=col_idx).data_type = "s"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


async def export_tenant_to_excel(db: AsyncSession, tenant_id: int) -> io.BytesIO:
    """
    Query operational tables for tenant_id and generate an in-memory .xlsx workbook.
    """
    # 1. Lookup maps for foreign IDs
    users_res = await db.execute(
        select(models.User.id, models.User.username, models.User.full_name)
        .where(models.User.tenant_id == tenant_id)
    )
    users_map: dict[int, str] = {
        row[0]: (row[2] or row[1]) for row in users_res.all()
    }

    patients_res = await db.execute(
        select(models.Patient.id, models.Patient.name)
        .where(models.Patient.tenant_id == tenant_id)
    )
    patients_map: dict[int, str] = {
        row[0]: row[1] for row in patients_res.all()
    }

    # 2. Patients
    patients_stmt = (
        select(models.Patient)
        .where(models.Patient.tenant_id == tenant_id, models.Patient.is_deleted == False)
        .order_by(models.Patient.id)
        .limit(ROW_LIMIT + 1)
    )
    patients_rows_db = (await db.execute(patients_stmt)).scalars().all()
    patients_headers = [
        "ID", "Name", "Phone", "Age", "Email", "Address", "Medical History", "Notes", "Created At"
    ]
    patients_data = [
        [
            p.id,
            p.name,
            p.phone,
            p.age,
            p.email,
            p.address,
            p.medical_history,
            p.notes,
            p.created_at,
        ]
        for p in patients_rows_db[:ROW_LIMIT]
    ]
    if len(patients_rows_db) > ROW_LIMIT:
        logger.warning("Patients sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        patients_data.append(TRUNCATION_MARKER)

    # 3. Appointments
    appt_stmt = (
        select(models.Appointment)
        .where(models.Appointment.tenant_id == tenant_id, models.Appointment.is_deleted == False)
        .order_by(models.Appointment.id)
        .limit(ROW_LIMIT + 1)
    )
    appt_rows_db = (await db.execute(appt_stmt)).scalars().all()
    appt_headers = [
        "ID", "Patient ID", "Patient Name", "Doctor ID", "Doctor Name", "Date Time", "Duration (mins)", "Status", "Notes"
    ]
    appt_data = [
        [
            a.id,
            a.patient_id,
            patients_map.get(a.patient_id, ""),
            a.doctor_id,
            users_map.get(a.doctor_id, ""),
            a.date_time,
            a.duration_minutes,
            a.status,
            a.notes,
        ]
        for a in appt_rows_db[:ROW_LIMIT]
    ]
    if len(appt_rows_db) > ROW_LIMIT:
        logger.warning("Appointments sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        appt_data.append(TRUNCATION_MARKER)

    # 4. Treatments
    treat_stmt = (
        select(models.Treatment)
        .where(models.Treatment.tenant_id == tenant_id, models.Treatment.is_deleted == False)
        .order_by(models.Treatment.id)
        .limit(ROW_LIMIT + 1)
    )
    treat_rows_db = (await db.execute(treat_stmt)).scalars().all()
    treat_headers = [
        "ID", "Patient ID", "Patient Name", "Doctor ID", "Doctor Name", "Procedure", "Tooth Number", "Cost", "Discount", "Status", "Date", "Notes"
    ]
    treat_data = [
        [
            t.id,
            t.patient_id,
            patients_map.get(t.patient_id, ""),
            t.doctor_id,
            users_map.get(t.doctor_id, ""),
            t.procedure,
            t.tooth_number,
            t.cost,
            t.discount,
            t.status,
            t.date,
            t.notes,
        ]
        for t in treat_rows_db[:ROW_LIMIT]
    ]
    if len(treat_rows_db) > ROW_LIMIT:
        logger.warning("Treatments sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        treat_data.append(TRUNCATION_MARKER)

    # 5. Payments
    pay_stmt = (
        select(models.Payment)
        .where(models.Payment.tenant_id == tenant_id)
        .order_by(models.Payment.id)
        .limit(ROW_LIMIT + 1)
    )
    pay_rows_db = (await db.execute(pay_stmt)).scalars().all()
    pay_headers = [
        "ID", "Patient ID", "Patient Name", "Doctor ID", "Doctor Name", "Amount", "Date", "Notes"
    ]
    pay_data = [
        [
            p.id,
            p.patient_id,
            patients_map.get(p.patient_id, ""),
            p.doctor_id,
            users_map.get(p.doctor_id, ""),
            p.amount,
            p.date,
            p.notes,
        ]
        for p in pay_rows_db[:ROW_LIMIT]
    ]
    if len(pay_rows_db) > ROW_LIMIT:
        logger.warning("Payments sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        pay_data.append(TRUNCATION_MARKER)

    # 6. Procedures
    proc_stmt = (
        select(models.Procedure)
        .where(models.Procedure.tenant_id == tenant_id)
        .order_by(models.Procedure.id)
        .limit(ROW_LIMIT + 1)
    )
    proc_rows_db = (await db.execute(proc_stmt)).scalars().all()
    proc_headers = ["ID", "Name", "Price"]
    proc_data = [
        [
            pr.id,
            pr.name,
            pr.price,
        ]
        for pr in proc_rows_db[:ROW_LIMIT]
    ]
    if len(proc_rows_db) > ROW_LIMIT:
        logger.warning("Procedures sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        proc_data.append(TRUNCATION_MARKER)

    # 7. Inventory Items
    stock_stmt = (
        select(models.StockItem)
        .options(
            joinedload(models.StockItem.warehouse),
            joinedload(models.StockItem.batch).joinedload(models.Batch.material),
        )
        .where(models.StockItem.tenant_id == tenant_id)
        .order_by(models.StockItem.id)
        .limit(ROW_LIMIT + 1)
    )
    stock_rows_db = (await db.execute(stock_stmt)).scalars().all()
    stock_headers = [
        "ID", "Warehouse ID", "Warehouse Name", "Material ID", "Material Name", "Batch Number", "Quantity", "Expiry Date"
    ]
    stock_data = [
        [
            s.id,
            s.warehouse_id,
            s.warehouse.name if s.warehouse else "",
            s.batch.material_id if s.batch else None,
            s.batch.material.name if s.batch and s.batch.material else "",
            s.batch.batch_number if s.batch else "",
            s.quantity,
            s.batch.expiry_date if s.batch else None,
        ]
        for s in stock_rows_db[:ROW_LIMIT]
    ]
    if len(stock_rows_db) > ROW_LIMIT:
        logger.warning("Inventory Items sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        stock_data.append(TRUNCATION_MARKER)

    # 8. Lab Orders
    lab_stmt = (
        select(models.LabOrder)
        .options(joinedload(models.LabOrder.laboratory))
        .where(models.LabOrder.tenant_id == tenant_id)
        .order_by(models.LabOrder.id)
        .limit(ROW_LIMIT + 1)
    )
    lab_rows_db = (await db.execute(lab_stmt)).scalars().all()
    lab_headers = [
        "ID", "Patient ID", "Patient Name", "Doctor ID", "Doctor Name", "Laboratory ID", "Laboratory Name", "Work Type", "Tooth Number", "Cost", "Price to Patient", "Status", "Order Date", "Delivery Date", "Notes"
    ]
    lab_data = [
        [
            lo.id,
            lo.patient_id,
            patients_map.get(lo.patient_id, ""),
            lo.doctor_id,
            users_map.get(lo.doctor_id, ""),
            lo.laboratory_id,
            lo.laboratory.name if lo.laboratory else "",
            lo.work_type,
            lo.tooth_number,
            lo.cost,
            lo.price_to_patient,
            lo.status,
            lo.order_date,
            lo.delivery_date,
            lo.notes,
        ]
        for lo in lab_rows_db[:ROW_LIMIT]
    ]
    if len(lab_rows_db) > ROW_LIMIT:
        logger.warning("Lab Orders sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        lab_data.append(TRUNCATION_MARKER)

    # 9. Treatment Plans
    plan_stmt = (
        select(models.ClinicalTreatmentPlan)
        .where(models.ClinicalTreatmentPlan.tenant_id == tenant_id)
        .order_by(models.ClinicalTreatmentPlan.id)
        .limit(ROW_LIMIT + 1)
    )
    plan_rows_db = (await db.execute(plan_stmt)).scalars().all()
    plan_headers = [
        "ID", "Patient ID", "Patient Name", "Title", "Status", "Created By User ID", "Created By User Name", "Created At", "Notes"
    ]
    plan_data = [
        [
            tp.id,
            tp.patient_id,
            patients_map.get(tp.patient_id, ""),
            tp.title,
            tp.status,
            tp.created_by_user_id,
            users_map.get(tp.created_by_user_id, ""),
            tp.created_at,
            tp.notes,
        ]
        for tp in plan_rows_db[:ROW_LIMIT]
    ]
    if len(plan_rows_db) > ROW_LIMIT:
        logger.warning("Treatment Plans sheet truncated at %d rows for tenant %s", ROW_LIMIT, tenant_id)
        plan_data.append(TRUNCATION_MARKER)

    sheets_data = {
        "Patients": (patients_headers, patients_data),
        "Appointments": (appt_headers, appt_data),
        "Treatments": (treat_headers, treat_data),
        "Payments": (pay_headers, pay_data),
        "Procedures": (proc_headers, proc_data),
        "Inventory Items": (stock_headers, stock_data),
        "Lab Orders": (lab_headers, lab_data),
        "Treatment Plans": (plan_headers, plan_data),
    }

    return await asyncio.to_thread(_build_workbook_sync, sheets_data)
